from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "production-make-ready-sequencing-blueprint.md"
OUTPUT = ROOT / "docs" / "production-make-ready-sequencing-blueprint.docx"

NAVY = "17324D"
TEAL = "167D8D"
PALE_BLUE = "EDF4F8"
PALE_GRAY = "F6F7F8"
MID_GRAY = "66717C"
LIGHT_BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def normalize_text(value: str) -> str:
    return (
        value.replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u00a0", " ")
    )


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str, *, color: RGBColor | None = None, size: float | None = None):
    text = normalize_text(text)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        set_run_font(run, "Aptos", size=size, bold=bold)
        if color:
            run.font.color.rgb = color


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), LIGHT_BORDER)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, "Aptos", 8.5)
    run.font.color.rgb = RGBColor(90, 100, 110)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(16)
    title_p_pr = title.element.get_or_add_pPr()
    title_p_bdr = title_p_pr.find(qn("w:pBdr"))
    if title_p_bdr is not None:
        title_p_pr.remove(title_p_bdr)

    for name, size, before, after in (
        ("Heading 1", 19, 20, 8),
        ("Heading 2", 14, 15, 5),
        ("Heading 3", 11.5, 10, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = False

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.1


def parse_table(lines: list[str], start: int):
    rows = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        row = [c.strip() for c in lines[idx].strip().strip("|").split("|")]
        rows.append(row)
        idx += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    return rows, idx


def add_table(doc: Document, rows: list[list[str]]):
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    repeat_table_header(table.rows[0])

    for r_idx, values in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = values[c_idx] if c_idx < len(values) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, value, color=RGBColor(255, 255, 255), size=9)
                for run in p.runs:
                    run.bold = True
            else:
                set_cell_shading(cell, PALE_BLUE if r_idx % 2 == 0 else "FFFFFF")
                numeric = bool(re.fullmatch(r"[\d$%.,+\- ]+", normalize_text(value)))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric and c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                add_inline(p, value, size=8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title_text = normalize_text(lines[0][2:].strip())
    section_titles = [normalize_text(line[2:].strip()) for line in lines if line.startswith("# ")][1:]

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True
    configure_styles(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(92)
    add_inline(p, title_text, size=30)
    remove_paragraph_borders(p)
    for run in p.runs:
        run.bold = True
        run.font.name = "Aptos Display"

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    add_inline(subtitle, "Business operating model and complete workflow map reference", color=RGBColor(50, 65, 78), size=15)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(22)
    meta.paragraph_format.line_spacing = 1.45
    add_inline(meta, "Audience  Business owners property management executives and operations leaders\n", color=RGBColor(70, 80, 90), size=10.5)
    add_inline(meta, f"Prepared  {date.today().strftime('%B %Y')}\n", color=RGBColor(70, 80, 90), size=10.5)
    add_inline(meta, "Scope  Vacant unit turnover from notice through rent ready handoff", color=RGBColor(70, 80, 90), size=10.5)

    tag = doc.add_paragraph()
    tag.paragraph_format.space_before = Pt(60)
    add_inline(tag, "The company keeps its systems and decision rights. The automation supplies the missing coordination memory between them.", color=RGBColor(20, 95, 108), size=12)
    tag.runs[0].italic = True

    doc.add_page_break()
    guide = doc.add_paragraph("Document guide", style="Heading 1")
    guide.paragraph_format.space_before = Pt(0)
    intro = doc.add_paragraph("The blueprint follows the sequence a decision-maker needs for evaluation and the sequence a diagrammer needs for a complete Excalidraw map.")
    for number, item in enumerate(section_titles, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(3)
        clean = re.sub(r"^\d+\s+", "", item)
        add_inline(p, f"{number}.  {clean}")
    doc.add_page_break()

    idx = 1
    first_body_h1 = True
    while idx < len(lines):
        raw = lines[idx]
        stripped = raw.strip()
        if not stripped:
            idx += 1
            continue
        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue
        if raw.startswith("# "):
            p = doc.add_paragraph(normalize_text(raw[2:].strip()), style="Heading 1")
            p.paragraph_format.space_before = Pt(0 if first_body_h1 else 20)
            p.paragraph_format.page_break_before = False
            first_body_h1 = False
        elif raw.startswith("## "):
            doc.add_paragraph(normalize_text(raw[3:].strip()), style="Heading 2")
        elif raw.startswith("### "):
            doc.add_paragraph(normalize_text(raw[4:].strip()), style="Heading 3")
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            add_inline(p, stripped)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)
        idx += 1

    for sec in doc.sections:
        sec.header.is_linked_to_previous = True
        header = sec.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.paragraph_format.space_after = Pt(0)
        run = header.add_run("Production Make Ready Sequencing Automation Blueprint")
        set_run_font(run, "Aptos", 8)
        run.font.color.rgb = RGBColor(100, 108, 116)

        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.paragraph_format.space_after = Pt(0)
        label = footer.add_run("Page ")
        set_run_font(label, "Aptos", 8.5)
        label.font.color.rgb = RGBColor(90, 100, 110)
        add_page_field(footer)

    props = doc.core_properties
    props.title = title_text
    props.subject = "Business operating model and workflow map reference for production make-ready sequencing automation"
    props.author = ""
    props.keywords = "property management, make-ready, turnover, automation, workflow, production blueprint"
    props.comments = ""

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
